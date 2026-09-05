"""
motor.py — Motor de reportes paramétrico.

Modos de sección (campo modo_bucle en reporte_seccion):
  N — Normal:   ejecuta SQL, toma la primera fila, coloca celdas en posición fija.
  F — Plano:    ejecuta SQL, itera sobre TODAS las filas avanzando paso_fila por iteración.
  G — Agrupado: ejecuta SQL, agrupa por campo_grupo, renderiza cabecera (tipo_celda='G')
                + ítems (tipo_celda='D') en una cuadrícula de num_columnas columnas.

Ancla (lk_seccion_ancla / ancla_gap):
  Cualquier modo puede anclar su posición a la última fila escrita por otra sección.
  delta = (ultima_fila_ancla + ancla_gap + 1) - min_fila_propia

Activación condicional (condicion_activa_sql):
  Si la consulta no devuelve filas, la sección se omite completamente.
"""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XlImage
from sqlalchemy import text
from cryptography.exceptions import InvalidTag
from sqlmodel import Session, select

from ..modelos.mis_modelos import ReporteCelda, ReporteImagen, ReporteImpreso, ReporteSeccion
from ..crypto import GesmedCrypto

_TEMPLATES = Path(__file__).parent.parent / "templates"

# ── Fechas ────────────────────────────────────────────────────────────────────

_MESES = ("enero","febrero","marzo","abril","mayo","junio",
          "julio","agosto","septiembre","octubre","noviembre","diciembre")

def _fmt_fecha(v: object) -> str:
    if isinstance(v, datetime):
        return f"{v.day} de {_MESES[v.month - 1]} {v.year}"
    if isinstance(v, date):
        return f"{v.day} de {_MESES[v.month - 1]} {v.year}"
    return str(v)

# ── Número a letras ───────────────────────────────────────────────────────────

_NL_LISTA = [
    "", "UNO", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE",
    "DIEZ", "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISEIS", "DIECISIETE",
    "DIECIOCHO", "DIECINUEVE", "VEINTE", "VEINTIUNO", "VEINTIDOS", "VEINTITRES",
    "VEINTICUATRO", "VEINTICINCO", "VEINTISEIS", "VEINTISIETE", "VEINTIOCHO", "VEINTINUEVE",
]
_NL_DECENAS  = ["", "DIEZ", "VEINTE", "TREINTA", "CUARENTA", "CINCUENTA",
                "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
_NL_CENTENAS = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS",
                "SEISCIENTOS", "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]

def _nl_menor_mil(n: int) -> str:
    if n < 30:   return _NL_LISTA[n]
    if n < 100:
        d, u = divmod(n, 10)
        return _NL_DECENAS[d] + ("" if u == 0 else " Y " + _NL_LISTA[u])
    if n == 100: return "CIEN"
    c, r = divmod(n, 100)
    base = _NL_CENTENAS[c]
    return base if r == 0 else base + " " + _nl_menor_mil(r)

def _entero_letras(n: int) -> str:
    if n == 0:   return "CERO"
    if n < 0:    return "MENOS " + _entero_letras(-n)
    if n < 1000: return _nl_menor_mil(n)
    miles, r = divmod(n, 1000)
    prefix = "MIL" if miles == 1 else _entero_letras(miles) + " MIL"
    return prefix if r == 0 else prefix + " " + _nl_menor_mil(r)

def _numero_a_letras(valor: str) -> str:
    s = valor.strip().replace(",", ".")
    try:
        f = float(s)
    except (ValueError, TypeError):
        return valor
    negativo = f < 0
    f        = abs(f)
    entero   = int(f)
    cents    = round((f - entero) * 100)
    texto    = _entero_letras(entero)
    if cents:
        texto += " CON " + _entero_letras(cents) + " CENTESIMOS"
    return ("MENOS " + texto) if negativo else texto

# ── Extracción de valor de una celda ─────────────────────────────────────────

def _colocar_titulo(ws, texto: str, nombre_fuente: str, fila: int, col: int) -> None:
    """Escribe el título de sección en la fila anterior al inicio del bucle."""
    if fila < 1:
        return
    cell = ws.cell(row=fila, column=col)
    if isinstance(cell, MergedCell):
        return
    cell.value = texto
    cell.font = Font(name=nombre_fuente, size=12, bold=True)

def _fmt_merge_v(celda: ReporteCelda) -> int:
    """Devuelve el merge_v configurado en el formato de la celda (mínimo 1)."""
    partes = celda.formato.split(".")
    if len(partes) >= 6:
        try:
            return max(1, int(partes[2]))
        except ValueError:
            pass
    return 1

def _get_valor(celda: ReporteCelda, resultado: dict, crypto) -> str:
    raw = resultado.get(celda.variable) if celda.variable else None
    if isinstance(raw, (bytes, bytearray)):
        if crypto:
            try:
                valor = crypto.desencriptar(bytes(raw))
            except (InvalidTag, ValueError) as e:
                print(f"[motor] desencriptación fallida — variable={celda.variable!r}, error={e}")
                valor = ""
        else:
            valor = ""
    elif raw is None:
        valor = ""
    elif isinstance(raw, (date, datetime)):
        valor = _fmt_fecha(raw)
    else:
        valor = str(raw)
    if celda.formato.endswith(".T") and valor:
        valor = _numero_a_letras(valor)
    return f"{celda.pre_fijo}{valor}{celda.post_fijo}"

# ── Punto de entrada público ──────────────────────────────────────────────────

def generar_reporte(nombre_reporte: str, contexto: dict[str, Any], engine) -> BytesIO:
    output, _ = _generar(nombre_reporte, contexto, engine)
    return output

def generar_reporte_con_avisos(
    nombre_reporte: str, contexto: dict[str, Any], engine
) -> tuple[BytesIO, list[str]]:
    return _generar(nombre_reporte, contexto, engine)

def generar_reporte_multisheet(
    nombre_reporte: str,
    base_ctx: dict[str, Any],
    tipos: list[tuple[int, str]],
    engine,
) -> tuple[BytesIO, list[str]]:
    """Genera un workbook con una hoja por cada (id_tipo, nombre_tipo) en tipos."""
    crypto = _init_crypto(base_ctx)
    avisos: list[str] = []
    with Session(engine) as session:
        impreso, plantilla = _cargar_impreso(nombre_reporte, session)
        wb       = load_workbook(str(plantilla))
        tpl_ws   = wb.active                          # hoja plantilla — se elimina al final
        imagenes = _cargar_imagenes(impreso, session)
        secciones = _cargar_secciones(impreso, session)
        for tipo_id, tipo_nombre in tipos:
            ws       = wb.copy_worksheet(tpl_ws)
            ws.title = tipo_nombre[:31]               # Excel: máx 31 caracteres
            ctx      = {**base_ctx, "lk_tipo": tipo_id}
            _renderizar_en_ws(ws, impreso, imagenes, secciones, ctx, crypto, session, avisos)
        wb.remove(tpl_ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, avisos

# ── Implementación interna ────────────────────────────────────────────────────

def _init_crypto(contexto: dict[str, Any]):
    if contexto.get("lk_medico"):
        try:
            return GesmedCrypto.para_medico(int(contexto["lk_medico"]))
        except Exception:
            pass
    return None

def _cargar_impreso(nombre_reporte: str, session):
    impreso = session.exec(
        select(ReporteImpreso).where(ReporteImpreso.nombre_reporte == nombre_reporte)
    ).first()
    if not impreso:
        raise ValueError(f"Reporte '{nombre_reporte}' no encontrado")
    plantilla = _TEMPLATES / f"{nombre_reporte}.xlsx"
    if not plantilla.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {plantilla}")
    return impreso, plantilla

def _cargar_imagenes(impreso, session):
    return session.exec(
        select(ReporteImagen)
        .where(ReporteImagen.lk_impreso == impreso.id_impreso)
        .order_by(ReporteImagen.orden)
    ).all()

def _cargar_secciones(impreso, session):
    return session.exec(
        select(ReporteSeccion)
        .where(ReporteSeccion.lk_impreso == impreso.id_impreso)
        .where(ReporteSeccion.es_activa == 1)
        .order_by(ReporteSeccion.orden)
    ).all()

def _renderizar_en_ws(ws, impreso, imagenes, secciones, contexto, crypto, session, avisos):
    """Renderiza el reporte completo en el worksheet ws."""
    # ── Imágenes ──────────────────────────────────────────────────────────────
    for _ri in imagenes:
        _img_path = _TEMPLATES / _ri.nombre_archivo
        if not _img_path.exists():
            continue
        _img        = XlImage(str(_img_path))
        _img.width  = _ri.ancho
        _img.height = _ri.alto
        _img.anchor = _ri.celda
        ws.add_image(_img)

    # ── Secciones ─────────────────────────────────────────────────────────────
    ultima_fila: dict[int, int] = {}

    for seccion in secciones:
        celdas = session.exec(
            select(ReporteCelda)
            .where(ReporteCelda.lk_seccion == seccion.id_seccion)
            .where(ReporteCelda.es_activa == 1)
        ).all()

        params_ctx = {
            k: contexto[k]
            for k in (seccion.parametros_in or [])
            if k in contexto
        }

        # ── Activación condicional ────────────────────────────────────────────
        cond_sql = (seccion.condicion_activa_sql or "").strip()
        if cond_sql and cond_sql.upper() != "NULL":
            fila_cond = session.execute(text(cond_sql), params_ctx).mappings().first()
            if not fila_cond:
                if celdas:
                    _delta = 0
                    if seccion.lk_seccion_ancla:
                        _min_f = min(c.fila for c in celdas)
                        _uf_a  = ultima_fila.get(seccion.lk_seccion_ancla, 0)
                        _delta = (_uf_a + seccion.ancla_gap + 1) - _min_f
                    ultima_fila[seccion.id_seccion] = max(
                        c.fila + _delta + _fmt_merge_v(c) - 1 for c in celdas
                    )
                else:
                    ultima_fila[seccion.id_seccion] = (
                        ultima_fila.get(seccion.lk_seccion_ancla, 0)
                        if seccion.lk_seccion_ancla else 0
                    )
                continue

        if not celdas:
            continue

        # ── Desplazamiento por ancla ──────────────────────────────────────────
        delta = 0
        if seccion.lk_seccion_ancla:
            min_fila_propia   = min(c.fila for c in celdas)
            ultima_fila_ancla = ultima_fila.get(seccion.lk_seccion_ancla, 0)
            delta = (ultima_fila_ancla + seccion.ancla_gap + 1) - min_fila_propia

        modo = seccion.modo_bucle or "N"

        # ── Modo N — normal ───────────────────────────────────────────────────
        if modo == "N":
            resultado: dict = {}
            if seccion.instruccion_sql:
                fila_sql = session.execute(
                    text(seccion.instruccion_sql), params_ctx
                ).mappings().first()
                if fila_sql:
                    resultado = dict(fila_sql)
            titulo = (seccion.titulo_seccion or "").strip()
            if len(titulo) >= 3:
                min_col    = min(c.columna for c in celdas)
                min_fila_n = min(c.fila for c in celdas)
                _colocar_titulo(ws, titulo, impreso.nombre_fuente,
                                min_fila_n + delta - 1, min_col)
            max_f = 0
            for celda in celdas:
                texto   = _get_valor(celda, resultado, crypto)
                fila_ef = celda.fila + delta
                _colocar(ws, celda, texto, impreso.nombre_fuente, fila_ef, celda.columna)
                max_f = max(max_f, fila_ef + _fmt_merge_v(celda) - 1)
            ultima_fila[seccion.id_seccion] = max_f

        # ── Modo F — bucle plano ──────────────────────────────────────────────
        elif modo == "F":
            filas_sql: list[dict] = []
            if seccion.instruccion_sql:
                filas_sql = [
                    dict(r) for r in session.execute(
                        text(seccion.instruccion_sql), params_ctx
                    ).mappings().all()
                ]
            n = len(filas_sql)
            if n > seccion.max_iteraciones:
                avisos.append(
                    f"Sección '{seccion.explica}': "
                    f"{n - seccion.max_iteraciones} ítem(s) truncados (máx {seccion.max_iteraciones})"
                )
                filas_sql = filas_sql[:seccion.max_iteraciones]
            titulo = (seccion.titulo_seccion or "").strip()
            if len(titulo) >= 3 and filas_sql:
                min_col    = min(c.columna for c in celdas)
                min_fila_f = min(c.fila for c in celdas)
                _colocar_titulo(ws, titulo, impreso.nombre_fuente,
                                min_fila_f + delta - 1, min_col)
            max_f = 0
            for i, resultado in enumerate(filas_sql):
                for celda in celdas:
                    texto   = _get_valor(celda, resultado, crypto)
                    fila_ef = celda.fila + delta + i * seccion.paso_fila
                    _colocar(ws, celda, texto, impreso.nombre_fuente, fila_ef, celda.columna)
                    max_f = max(max_f, fila_ef + _fmt_merge_v(celda) - 1)
            ultima_fila[seccion.id_seccion] = max_f

        # ── Modo G — bucle agrupado ───────────────────────────────────────────
        elif modo == "G":
            filas_sql = []
            if seccion.instruccion_sql:
                filas_sql = [
                    dict(r) for r in session.execute(
                        text(seccion.instruccion_sql), params_ctx
                    ).mappings().all()
                ]
            campo_g   = seccion.campo_grupo or ""
            num_cols  = max(1, seccion.num_columnas)
            paso_col  = seccion.paso_columna
            slot_h    = max(1, seccion.slot_height)
            paso_fila = max(1, seccion.paso_fila)
            celdas_G  = [c for c in celdas if c.tipo_celda == "G"]
            celdas_D  = [c for c in celdas if c.tipo_celda != "G"]
            min_fila  = min(c.fila for c in celdas)
            grupos: dict[str, list[dict]] = defaultdict(list)
            for fila in filas_sql:
                clave = str(fila.get(campo_g, "")) if campo_g else ""
                grupos[clave].append(fila)
            max_row_bucket = (len(grupos) - 1) // num_cols if grupos else 0
            max_f  = 0
            titulo = (seccion.titulo_seccion or "").strip()
            if len(titulo) >= 3 and grupos:
                min_col = min(c.columna for c in celdas)
                _colocar_titulo(ws, titulo, impreso.nombre_fuente,
                                min_fila + delta - 1, min_col)
            for i, (nombre_grupo, items) in enumerate(grupos.items()):
                col_bucket     = i % num_cols
                row_bucket     = i // num_cols
                es_ultima_fila = (row_bucket == max_row_bucket)
                col_offset     = col_bucket * paso_col
                fila_slot      = min_fila + delta + row_bucket * slot_h
                for celda in celdas_G:
                    fila_ef = fila_slot + (celda.fila - min_fila)
                    col_ef  = celda.columna + col_offset
                    _colocar(ws, celda, nombre_grupo, impreso.nombre_fuente, fila_ef, col_ef)
                    max_f = max(max_f, fila_ef + _fmt_merge_v(celda) - 1)
                max_items = len(items) if es_ultima_fila else seccion.max_iteraciones
                if not es_ultima_fila and len(items) > seccion.max_iteraciones:
                    avisos.append(
                        f"Grupo '{nombre_grupo}': "
                        f"{len(items) - seccion.max_iteraciones} ítem(s) truncados"
                    )
                n_insertadas = len(items[:max_items])
                for j, resultado in enumerate(items[:max_items]):
                    for celda in celdas_D:
                        texto   = _get_valor(celda, resultado, crypto)
                        fila_ef = fila_slot + (celda.fila - min_fila) + j * paso_fila
                        col_ef  = celda.columna + col_offset
                        _colocar(ws, celda, texto, impreso.nombre_fuente, fila_ef, col_ef)
                        max_f = max(max_f, fila_ef + _fmt_merge_v(celda) - 1)
                if n_insertadas < slot_h - 1 and celdas_D:
                    celda_ref = celdas_D[0]
                    fila_ef   = fila_slot + (celda_ref.fila - min_fila) + n_insertadas * paso_fila
                    col_ef    = celda_ref.columna + col_offset
                    _colocar(ws, celda_ref, "- - - -", impreso.nombre_fuente, fila_ef, col_ef)
            ultima_fila[seccion.id_seccion] = max_f

def _generar(
    nombre_reporte: str, contexto: dict[str, Any], engine
) -> tuple[BytesIO, list[str]]:
    crypto  = _init_crypto(contexto)
    avisos: list[str] = []
    with Session(engine) as session:
        impreso, plantilla = _cargar_impreso(nombre_reporte, session)
        wb        = load_workbook(str(plantilla))
        ws        = wb.active
        imagenes  = _cargar_imagenes(impreso, session)
        secciones = _cargar_secciones(impreso, session)
        _renderizar_en_ws(ws, impreso, imagenes, secciones, contexto, crypto, session, avisos)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, avisos

# ── Colocación de celda ───────────────────────────────────────────────────────

def _colocar(
    ws,
    celda: ReporteCelda,
    texto: str,
    nombre_fuente: str,
    fila: int,
    col: int,
) -> None:
    """Aplica valor, fuente, merge, alineación, sombra y bordes en (fila, col)."""
    partes = celda.formato.split(".")
    if len(partes) not in (6, 7):
        raw = ws.cell(row=fila, column=col)
        if not isinstance(raw, MergedCell):
            raw.value = texto
        return

    size     = int(partes[0])
    merge_h  = int(partes[1])
    merge_v  = int(partes[2])
    sombra   = int(partes[3])
    efecto   = partes[4][0]
    alinea_h = partes[4][1]
    alinea_v = partes[4][2]
    borders  = partes[5]

    if isinstance(ws.cell(row=fila, column=col), MergedCell):
        return

    def _side(flag: str) -> Side:
        return Side(border_style="thin" if flag == "1" else None)

    # Para celdas combinadas los bordes derecho/inferior visibles pertenecen a las
    # celdas extremas del rango (no al anchor). Deben aplicarse ANTES de merge_cells(),
    # porque después esas celdas se convierten en MergedCell y openpyxl no las deja editar.
    if len(borders) == 4 and (merge_h > 1 or merge_v > 1):
        last_col = col  + merge_h - 1
        last_row = fila + merge_v - 1
        if borders[1] == "1" and merge_h > 1:
            rc = ws.cell(row=fila, column=last_col)
            if not isinstance(rc, MergedCell):
                rc.border = Border(right=Side(border_style="thin"))
        if borders[2] == "1" and merge_v > 1:
            bc = ws.cell(row=last_row, column=col)
            if not isinstance(bc, MergedCell):
                bc.border = Border(bottom=Side(border_style="thin"))

    if merge_h > 1 or merge_v > 1:
        ws.merge_cells(
            start_row=fila,   start_column=col,
            end_row=fila + merge_v - 1,
            end_column=col  + merge_h - 1,
        )

    cell = ws.cell(row=fila, column=col)
    cell.value = texto
    cell.font  = Font(
        name=nombre_fuente, size=size,
        bold=(efecto == "B"),
        italic=(efecto == "Q"),
        underline=("single" if efecto == "S" else None),
    )
    _H = {"L": "left", "C": "center", "R": "right"}
    _V = {"U": "top",  "C": "center", "D": "bottom"}
    cell.alignment = Alignment(
        horizontal=_H.get(alinea_h, "left"),
        vertical=_V.get(alinea_v, "top"),
        wrap_text=(merge_h > 1 or merge_v > 1),
    )
    if sombra > 0:
        g = int((1 - sombra / 100) * 255)
        cell.fill = PatternFill("solid", fgColor=f"{g:02X}{g:02X}{g:02X}")

    if len(borders) == 4:
        cell.border = Border(
            top=_side(borders[0]),
            right=_side(borders[1]) if merge_h == 1 else Side(),
            bottom=_side(borders[2]) if merge_v == 1 else Side(),
            left=_side(borders[3]),
        )
